"""Sample data for trying the Data & Reporting agent: a small Indian D2C
snacks brand, April to 11 October 2026. Seeded, so it's the same every run."""
import csv
import random
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font

out = sys.argv[1]
rng = random.Random(42)
START, END = date(2026, 4, 1), date(2026, 10, 11)

PRODUCTS = [
    # name, category, price
    ("Masala Makhana 100g", "Snacks", 149), ("Peri Peri Makhana 100g", "Snacks", 149),
    ("Roasted Chana Mix 200g", "Snacks", 129), ("Ragi Chips 90g", "Snacks", 99),
    ("Jaggery Peanut Chikki (12)", "Sweets", 199), ("Dry Fruit Laddoo Box", "Sweets", 549),
    ("Kaju Katli 250g", "Sweets", 449), ("Assam Masala Chai 250g", "Beverages", 299),
    ("Kashmiri Kahwa 100g", "Beverages", 349), ("Filter Coffee Powder 250g", "Beverages", 279),
    ("Festive Hamper Small", "Gift Boxes", 999), ("Festive Hamper Large", "Gift Boxes", 1899),
]
CITIES = [("Mumbai", "Maharashtra", 16), ("Pune", "Maharashtra", 9), ("Bengaluru", "Karnataka", 14), ("Delhi", "Delhi", 13),
          ("Gurugram", "Haryana", 6), ("Hyderabad", "Telangana", 9), ("Chennai", "Tamil Nadu", 8), ("Ahmedabad", "Gujarat", 7),
          ("Kolkata", "West Bengal", 6), ("Jaipur", "Rajasthan", 4), ("Lucknow", "Uttar Pradesh", 4), ("Kochi", "Kerala", 4)]
CHANNELS = [("Website", 34), ("Amazon", 28), ("Flipkart", 14), ("Instagram Shop", 12), ("Offline store", 12)]
PAYMENTS = [("UPI", 52), ("Card", 18), ("COD", 24), ("NetBanking", 6)]
FIRST = ["Aarav", "Diya", "Ishaan", "Ananya", "Rohan", "Meera", "Kabir", "Saanvi", "Arjun", "Priya", "Vihaan", "Neha",
         "Aditya", "Kavya", "Rahul", "Sneha", "Karan", "Pooja", "Siddharth", "Riya", "Manish", "Tanvi", "Nikhil", "Shreya"]
LAST = ["Sharma", "Iyer", "Patel", "Reddy", "Nair", "Gupta", "Mehta", "Singh", "Kulkarni", "Das", "Bose", "Joshi", "Rao", "Khan", "Verma"]


def pick(options):
    return rng.choices([o[0] for o in options], weights=[o[-1] for o in options])[0]


def rupees(n):
    """₹1,24,500 (Indian grouping), as people type it in sheets."""
    s = str(int(round(n)))
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        s = ",".join(parts) + "," + tail
    return f"₹{s}"


# --- orders -------------------------------------------------------------------
returning: list[str] = []
orders = []
day, n = START, 10001
while day <= END:
    # Growth into the festive season, more on weekends.
    base = 8 + (day - START).days / 22
    if day.month == 10:
        base *= 1.6
    if day.weekday() >= 5:
        base *= 1.3
    for _ in range(max(0, int(rng.gauss(base, 1.1)))):
        product, category, price = rng.choice(PRODUCTS[:10] if day.month < 9 else PRODUCTS)
        qty = rng.choices([1, 2, 3, 4, 6], weights=[55, 25, 10, 6, 4])[0]
        discount = rng.choices([0, 5, 10, 15, 20], weights=[45, 20, 20, 10, 5])[0]
        channel, payment = pick(CHANNELS), pick(PAYMENTS)
        if channel == "Offline store" and payment == "COD":
            payment = "UPI"
        city, state, _ = rng.choices(CITIES, weights=[c[2] for c in CITIES])[0]
        # Repeat orders build up as the brand finds its regulars.
        repeat = bool(returning) and rng.random() < 0.12 + 0.25 * (day - START).days / (END - START).days
        customer = rng.choice(returning) if repeat else f"{rng.choice(FIRST)} {rng.choice(LAST)}"
        if not repeat:
            returning.append(customer)
        returned_odds = 0.11 if payment == "COD" else 0.03
        status = rng.choices(["Delivered", "Returned", "Cancelled"], weights=[1 - returned_odds - 0.03, returned_odds, 0.03])[0]
        if (END - day).days < 4 and status == "Delivered":
            status = "Shipped"
        total = price * qty * (100 - discount) / 100
        orders.append([f"ORD-{n}", day.strftime("%d/%m/%Y"), customer, city, state, channel, product, category, qty,
                       rupees(price), f"{discount}%", rupees(total), payment, status, "Yes" if repeat else "No"])
        n += 1
    day += timedelta(days=1)

with open(f"{out}/sales_orders.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["Order ID", "Order Date", "Customer", "City", "State", "Channel", "Product", "Category", "Quantity",
                "Unit Price", "Discount", "Order Total", "Payment Method", "Status", "Repeat Customer"])
    w.writerows(orders)

# --- expenses -----------------------------------------------------------------
expenses = []
month = date(2026, 4, 1)
while month <= END:
    m = month.month
    festive = 1.8 if m == 10 else 1.3 if m == 9 else 1.0
    rows = [
        (1, "Salaries", "Team payroll", 120000 + (m - 4) * 5000, "Bank transfer", "No"),
        (3, "Rent", "Warehouse, Bhiwandi", 35000, "Bank transfer", "Yes"),
        (5, "Software", "Shopify", 7400, "Card", "Yes"),
        (5, "Software", "Zoho One", 4300, "Card", "Yes"),
        (7, "Utilities", "Electricity and internet", rng.randint(9000, 13000), "UPI", "Yes"),
    ]
    for week in (2, 9, 16, 23):
        rows.append((week, "Ads - Meta", "Meta Ads", rng.randint(8000, 14000) * festive, "Card", "Yes"))
        rows.append((week + 1, "Ads - Google", "Google Ads", rng.randint(5000, 9000) * festive, "Card", "Yes"))
        rows.append((week + 2, "Packaging", rng.choice(["BoxKraft Pvt Ltd", "PackRight India"]), rng.randint(4000, 7000) * festive, "Bank transfer", "Yes"))
        rows.append((week + 3, "Shipping", rng.choice(["Delhivery", "Shiprocket", "Blue Dart"]), rng.randint(6000, 10000) * festive, "Bank transfer", "Yes"))
    if m == 8:
        rows.append((20, "Marketing", "Influencer campaign (Rakhi)", 45000, "Bank transfer", "No"))
    if m == 10:
        rows.append((4, "Marketing", "Diwali influencer campaign", 80000, "Bank transfer", "No"))
    for d, category, vendor, amount, paid_by, gst in rows:
        when = date(month.year, m, min(d, 28))
        if when <= END:
            expenses.append([when.isoformat(), category, vendor, rupees(amount), paid_by, gst])
    month = date(2026, m + 1, 1) if m < 12 else date(2027, 1, 1)
expenses.sort(key=lambda r: r[0])
with open(f"{out}/expenses.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["Date", "Category", "Vendor", "Amount", "Paid By", "GST Included"])
    w.writerows(expenses)

# --- marketing campaigns (Excel) ----------------------------------------------
campaigns = [
    ("Summer Makhana Launch", "Meta Ads", date(2026, 4, 5), date(2026, 4, 30)),
    ("Chai Lovers Search", "Google Ads", date(2026, 4, 10), date(2026, 5, 31)),
    ("Monsoon Chai Combo", "Meta Ads", date(2026, 6, 1), date(2026, 6, 30)),
    ("Healthy Snacking Reels", "Instagram", date(2026, 6, 15), date(2026, 7, 15)),
    ("Amazon Sponsored Products", "Amazon Ads", date(2026, 5, 1), date(2026, 10, 11)),
    ("Rakhi Gift Boxes", "Meta Ads", date(2026, 8, 1), date(2026, 8, 19)),
    ("Rakhi Influencers", "Influencers", date(2026, 8, 5), date(2026, 8, 19)),
    ("Independence Day Sale", "Google Ads", date(2026, 8, 10), date(2026, 8, 16)),
    ("Back to Office Snacks", "LinkedIn Ads", date(2026, 9, 1), date(2026, 9, 30)),
    ("Ganesh Chaturthi Sweets", "Meta Ads", date(2026, 9, 10), date(2026, 9, 20)),
    ("Diwali Hampers Early Bird", "Meta Ads", date(2026, 10, 1), date(2026, 10, 11)),
    ("Diwali Hampers Search", "Google Ads", date(2026, 10, 1), date(2026, 10, 11)),
    ("Diwali Influencers", "Influencers", date(2026, 10, 4), date(2026, 10, 11)),
    ("Festive WhatsApp Broadcast", "WhatsApp", date(2026, 10, 6), date(2026, 10, 11)),
]
wb = Workbook()
ws = wb.active
ws.title = "Campaigns"
header = ["Campaign", "Channel", "Start Date", "End Date", "Spend (₹)", "Impressions", "Clicks", "Leads", "Orders", "Revenue (₹)", "ROAS"]
ws.append(header)
# channel: (cost per 1,000 views ₹, cost per click ₹, click-to-order rate)
COSTS = {"Meta Ads": (110, 9, 0.028), "Google Ads": (260, 14, 0.04), "Instagram": (90, 8, 0.022),
         "Amazon Ads": (300, 12, 0.07), "Influencers": (45, 18, 0.02), "LinkedIn Ads": (420, 65, 0.012), "WhatsApp": (60, 2.5, 0.09)}
for name, channel, start, end in campaigns:
    days = (end - start).days + 1
    spend = rng.uniform(700, 1500) * days * (1.4 if "Diwali" in name else 1)
    if channel == "WhatsApp":
        spend = rng.uniform(3000, 5000)
    if channel == "Influencers":
        spend = rng.uniform(40000, 80000)
    cpm, cpc, conversion = COSTS[channel]
    impressions = int(spend / cpm * 1000 * rng.uniform(0.85, 1.15))
    clicks = int(spend / cpc * rng.uniform(0.85, 1.15))
    orders_n = max(1, int(clicks * conversion * rng.uniform(0.8, 1.25)))
    aov = rng.uniform(430, 560) * (2.4 if "Hamper" in name or "Gift" in name else 1)
    revenue = orders_n * aov
    ws.append([name, channel, start, end, round(spend), impressions, clicks, int(clicks * rng.uniform(0.04, 0.09)), orders_n,
               round(revenue), round(revenue / spend, 2)])
for cell in ws[1]:
    cell.font = Font(bold=True)
for row in ws.iter_rows(min_row=2):
    row[2].number_format = row[3].number_format = "DD-MMM-YYYY"
    row[4].number_format = row[9].number_format = '"₹"#,##0'
for col, width in zip("ABCDEFGHIJK", (30, 14, 13, 13, 11, 13, 9, 8, 8, 12, 7)):
    ws.column_dimensions[col].width = width
wb.save(f"{out}/marketing_campaigns.xlsx")

# --- inventory ----------------------------------------------------------------
with open(f"{out}/inventory.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["SKU", "Product", "Category", "Stock On Hand", "Reorder Level", "Unit Cost", "Supplier", "Last Restocked"])
    suppliers = {"Snacks": "Nashik Agro Foods", "Sweets": "Rajkot Mithai Co", "Beverages": "Nilgiri Estates", "Gift Boxes": "In-house packing"}
    for i, (product, category, price) in enumerate(PRODUCTS, 1):
        stock = rng.choice([18, 35, 60, 120, 240, 310, 75, 12, 95, 400, 44, 150])
        w.writerow([f"SKU-{100 + i}", product, category, stock, rng.choice([40, 50, 80, 100]), rupees(price * rng.uniform(0.42, 0.58)),
                    suppliers[category], (END - timedelta(days=rng.randint(3, 60))).strftime("%d %b %Y")])

print(len(orders), "orders,", len(expenses), "expenses,", len(campaigns), "campaigns,", len(PRODUCTS), "products")
